"""Read-only local ingestion with SQLite, provenance and minimized excerpts."""
from __future__ import annotations

import csv
import hashlib
import importlib.util
import io
import json
import re
import sqlite3
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path

from .config import LIBRARY_TYPES, validate_settings

TEXT_FORMATS = {'.txt', '.md', '.csv', '.json', '.jsonl', '.html', '.htm'}
IMAGE_FORMATS = {'.png', '.jpg', '.jpeg', '.webp'}
OPTIONAL_FORMATS = {'.docx': 'docx', '.xlsx': 'openpyxl', '.pdf': 'pypdf'}


def redact(text: str) -> str:
    """Deterministic local minimization; this is not a complete semantic DLP system."""
    value = str(text)
    value = re.sub(r'(?<!\d)(?:\+?86[- ]?)?1[3-9]\d{9}(?!\d)', '[电话已脱敏]', value)
    value = re.sub(r'(?<![A-Za-z0-9._%+-])[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}(?![A-Za-z])', '[邮箱已脱敏]', value)
    value = re.sub(r'(?<!\d)\d{17}[\dXx](?!\d)', '[证件已脱敏]', value)
    value = re.sub(r'(?:姓名|联系人|客户姓名|收件人|微信|微信号|订单号|订单编号|地址|住址|收货地址)\s*[:：]\s*[^\s,，;；。]+', '[个人信息已脱敏]', value)
    value = re.sub(r'(?:我叫|本人叫)[\u4e00-\u9fff]{2,4}', '[姓名已脱敏]', value)
    value = re.sub(r'(?:病史|诊断|身份证|健康详情)\s*[:：]\s*[^\n;；。]+', '[个人健康信息已脱敏]', value)
    value = re.sub(r'\b(?:sk-[A-Za-z0-9_-]{16,}|Bearer\s+\S+)', '[凭据已脱敏]', value, flags=re.I)
    return value


def _safe_metadata(value):
    if isinstance(value, dict):
        return {str(k): _safe_metadata(v) for k, v in value.items()
                if not re.search(r'password|secret|api.?key|token|phone|email|address|customer_name|order_id', str(k), re.I)}
    if isinstance(value, list):
        return [_safe_metadata(v) for v in value]
    return redact(value) if isinstance(value, str) else value


def _decode(data: bytes):
    for encoding in ('utf-8-sig', 'gb18030', 'big5'):
        try:
            text = data.decode(encoding, errors='strict')
            if '\x00' in text:
                raise ValueError('文本含 NUL，无法完整解析')
            return text, 'utf-8-bom' if encoding == 'utf-8-sig' and data.startswith(b'\xef\xbb\xbf') else ('utf-8' if encoding == 'utf-8-sig' else encoding)
        except UnicodeError:
            continue
    raise ValueError('decode_failed: UTF-8/GB18030/Big5 均无法完整解码')


class _BodyParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.skip = 0
        self.parts = []

    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style', 'noscript'):
            self.skip += 1

    def handle_endtag(self, tag):
        if tag in ('script', 'style', 'noscript'):
            self.skip = max(0, self.skip - 1)

    def handle_data(self, data):
        if not self.skip and data.strip():
            self.parts.append((data.strip(), self.getpos()[0]))


def _role(value):
    v = str(value or '').strip().lower()
    if v in ('customer', 'client', 'user', '客户', '买家', '用户'):
        return 'customer'
    if v in ('service', 'assistant', '客服', '商家', '销售'):
        return 'service'
    return 'unknown'


def _structured_record(row, line, location_extra=None):
    if isinstance(row, dict):
        text = next((row[k] for k in ('text', 'content', 'message', '正文', '内容') if k in row), None)
        if text is None:
            excluded = {'role', 'date', 'timestamp', 'conversation_id', 'session_id'}
            text = '\n'.join(f'{k}：{v}' for k, v in row.items() if k not in excluded)
        meta = {'role': _role(row.get('role', row.get('角色'))),
                'date': row.get('date', row.get('timestamp', row.get('日期'))),
                'conversation_id': row.get('conversation_id', row.get('session_id', row.get('会话标识')))}
    else:
        text = str(row)
        meta = {'role': 'unknown', 'date': None, 'conversation_id': None}
    role_text = str(text)
    date_prefix = re.match(r'^\s*\[?(\d{4}-\d{2}-\d{2})(?:[ T]\d{2}:\d{2}(?::\d{2})?)?\]?\s+', role_text)
    if date_prefix:
        try:
            datetime.strptime(date_prefix.group(1), '%Y-%m-%d')
            if meta['date'] is None:
                meta['date'] = date_prefix.group(1)
            role_text = role_text[date_prefix.end():]
        except ValueError:
            pass
    if meta['role'] == 'unknown':
        match = re.match(r'^\s*(客户|买家|用户|客服|商家|销售|customer|service)\s*[:：]', role_text, re.I)
        if match:
            meta['role'] = _role(match.group(1))
    return str(text), {'line_start': line, 'line_end': line, **(location_extra or {})}, _safe_metadata(meta)


def _parse_text(data, suffix):
    text, encoding = _decode(data)
    result = []
    if suffix == '.json':
        parsed = json.loads(text)
        rows = parsed if isinstance(parsed, list) else parsed.get('messages', [parsed]) if isinstance(parsed, dict) else [parsed]
        if not isinstance(rows, list):
            raise ValueError('messages 必须是数组')
        for i, row in enumerate(rows):
            result.append(_structured_record(row, 1, {'json_pointer': f'/{i}' if isinstance(parsed, list) else f'/messages/{i}' if isinstance(parsed, dict) and 'messages' in parsed else ''}))
    elif suffix == '.jsonl':
        for number, line in enumerate(text.splitlines(), 1):
            if line.strip():
                result.append(_structured_record(json.loads(line), number))
    elif suffix == '.csv':
        reader = csv.DictReader(io.StringIO(text))
        for row in reader:
            result.append(_structured_record(row, reader.line_num, {'csv_row': reader.line_num}))
    elif suffix in ('.html', '.htm'):
        parser = _BodyParser()
        parser.feed(text)
        result = [_structured_record(part, line) for part, line in parser.parts]
    else:
        result = [_structured_record(line, number) for number, line in enumerate(text.splitlines(), 1) if line.strip()]
    return result, encoding


def parser_status():
    return {**{extension: {'available': True, 'parser': 'stdlib'} for extension in TEXT_FORMATS},
            **{extension: {'available': importlib.util.find_spec(module) is not None, 'parser': module}
               for extension, module in {**OPTIONAL_FORMATS, **dict.fromkeys(IMAGE_FORMATS, 'PIL')}.items()}}


def _optional_parse(path, suffix):
    module = OPTIONAL_FORMATS[suffix]
    if importlib.util.find_spec(module) is None:
        raise ImportError(f'optional_parser_unavailable:{module}')
    result = []
    if suffix == '.docx':
        from docx import Document
        document = Document(path)
        for number, paragraph in enumerate(document.paragraphs, 1):
            if paragraph.text.strip():
                result.append(_structured_record(paragraph.text, number, {'paragraph': number}))
        for table_number, table in enumerate(document.tables, 1):
            for row_number, row in enumerate(table.rows, 1):
                result.append(_structured_record(' | '.join(c.text for c in row.cells), row_number,
                                                 {'table': table_number, 'row': row_number}))
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook:
                for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    if any(v is not None for v in row):
                        result.append(_structured_record(' | '.join(str(v or '') for v in row), number,
                                                         {'sheet': sheet.title, 'row': number}))
        finally:
            workbook.close()
    else:
        from pypdf import PdfReader
        reader = PdfReader(path)
        for number, page in enumerate(reader.pages, 1):
            body = page.extract_text() or ''
            if body.strip():
                result.append(_structured_record(body, 1, {'page': number}))
        if not result:
            raise ValueError('pdf_no_extractable_text: 需要实际 OCR 能力，未识别')
    return result, 'parser:' + module


def _ngrams(value):
    normalized = ''.join(re.findall(r'[\w\u4e00-\u9fff]+', value.lower()))
    return set(normalized) | {normalized[i:i+2] for i in range(len(normalized)-1)}


class LibraryIndex:
    def __init__(self, settings):
        self.settings = validate_settings(settings)
        if not self.settings.get('workspace_root'):
            raise ValueError('缺少 workspace_root')
        self.workspace = Path(self.settings['workspace_root'])
        self.workspace.mkdir(parents=True, exist_ok=True)
        self.db_path = self.workspace / 'index.sqlite'
        if self.db_path.is_symlink():
            raise ValueError('索引数据库不能为符号链接')
        with self._connect() as db:
            db.executescript('''
                CREATE TABLE IF NOT EXISTS files (
                    file_key TEXT PRIMARY KEY, library_type TEXT NOT NULL,
                    relative_path TEXT NOT NULL, hash TEXT, fingerprint TEXT,
                    status TEXT NOT NULL, reason TEXT);
                CREATE TABLE IF NOT EXISTS sources (
                    source_id TEXT PRIMARY KEY, file_key TEXT NOT NULL,
                    library_type TEXT NOT NULL, product_id TEXT NOT NULL,
                    hash TEXT NOT NULL, location TEXT NOT NULL,
                    snippet TEXT NOT NULL, metadata TEXT NOT NULL);
                CREATE INDEX IF NOT EXISTS source_filter ON sources(library_type, product_id);
            ''')

    def _connect(self):
        connection = sqlite3.connect(self.db_path, timeout=30)
        connection.row_factory = sqlite3.Row
        return connection

    def _product(self, file_key, source_id):
        product_id = self.settings.get('source_mappings', {}).get(file_key)
        matches = []
        for product in self.settings['products']:
            if source_id in product.get('source_ids', []) + product.get('image_source_ids', []):
                matches.append(product['product_id'])
            if file_key in product.get('file_mappings', {}):
                matches.append(product['product_id'])
        if product_id:
            matches.append(product_id)
        if len(set(matches)) > 1:
            raise ValueError('product_mapping_conflict')
        if matches:
            product = next(p for p in self.settings['products'] if p['product_id'] == matches[0])
            return product['product_id'], product['version']
        return 'general', None

    def _sidecar(self, path, root):
        combined, hashes = {}, []
        for extension in ('.txt', '.md', '.json'):
            candidate = path.with_suffix(extension)
            if candidate.exists():
                actual = candidate.resolve()
                if root not in actual.parents:
                    raise ValueError('sidecar_outside_library')
                raw = candidate.read_bytes()
                hashes.append(hashlib.sha256(raw).hexdigest())
                body, _ = _decode(raw)
                if extension == '.json':
                    value = json.loads(body)
                    if not isinstance(value, dict):
                        raise ValueError('image_sidecar_requires_object')
                    combined.update(value)
                else:
                    combined[extension[1:]] = body
        return _safe_metadata(combined), hashes

    def _parse(self, path, root, raw):
        suffix = path.suffix.lower()
        if suffix in TEXT_FORMATS:
            return _parse_text(raw, suffix)
        if suffix in OPTIONAL_FORMATS:
            return _optional_parse(path, suffix)
        if suffix in IMAGE_FORMATS:
            if importlib.util.find_spec('PIL') is None:
                raise ImportError('optional_parser_unavailable:Pillow')
            from PIL import Image
            with Image.open(io.BytesIO(raw)) as opened:
                fmt, size = opened.format, opened.size
                opened.verify()
            with Image.open(io.BytesIO(raw)) as opened:
                opened.load()
            allowed = {'.png': 'PNG', '.jpg': 'JPEG', '.jpeg': 'JPEG', '.webp': 'WEBP'}
            if fmt != allowed[suffix]:
                raise ValueError('image_extension_mismatch')
            sidecar, _ = self._sidecar(path, root)
            meta = {'format': fmt, 'width': size[0], 'height': size[1], 'is_image': True,
                    'sidecar': sidecar, 'sidecar_trust': 'untrusted', 'ocr_performed': False,
                    'allowed_reference_factors': sidecar.get('allowed_reference_factors', []),
                    'forbidden_inheritance': sidecar.get('forbidden_inheritance', ['旧型号', '参数', '文字', 'Logo', '功能暗示'])}
            for key in ('approved', 'status', 'external_use_approved', 'version', 'immutable', 'borrow', 'forbidden', 'description'):
                if key in sidecar:
                    meta[key] = sidecar[key]
            return [(path.stem + ' ' + json.dumps(sidecar, ensure_ascii=False), {'line_start': 1, 'line_end': 1}, meta)], 'binary:image'
        raise ValueError('unsupported_format')

    def update(self, library_type=None):
        if library_type is not None and library_type not in LIBRARY_TYPES:
            raise ValueError('未知资料库类型')
        # Validate roots again on each scan to catch changed symlink targets.
        self.settings = validate_settings(self.settings)
        report = dict(added=0, modified=0, deleted=0, unchanged=0, parsed=0, failed=0, excluded=0,
                      files_seen=0, issues=[], by_library={}, coverage_basis='configured_files',
                      deduplication='No customer count; source_id identifies a file hash and location. Conversation IDs only when explicit.')
        with self._connect() as db:
            for kind in LIBRARY_TYPES:
                if library_type and library_type != kind:
                    continue
                root_value = self.settings['libraries'].get(kind)
                if root_value is None:
                    report['by_library'][kind] = {'configured': False, 'parsed': 0, 'failed': 0, 'excluded': 0}
                    continue
                root = Path(root_value)
                before = {r['file_key']: dict(r) for r in db.execute('SELECT * FROM files WHERE library_type=?', (kind,))}
                seen = set()
                counts = dict(configured=True, parsed=0, failed=0, excluded=0)
                # rglob does not recurse into directory symlinks on supported Python 3.11.
                for path in sorted(root.rglob('*')):
                    if path.is_dir() and not path.is_symlink():
                        continue
                    relative = path.relative_to(root).as_posix()
                    file_key = kind + '/' + relative
                    seen.add(file_key)
                    report['files_seen'] += 1
                    status, reason, digest, fingerprint = 'parsed', None, None, None
                    try:
                        actual = path.resolve()
                        if root not in actual.parents:
                            raise PermissionError('path_outside_library')
                        if not actual.is_file():
                            raise PermissionError('symlink_directory_excluded')
                        suffix = path.suffix.lower()
                        if suffix not in TEXT_FORMATS | IMAGE_FORMATS | set(OPTIONAL_FORMATS):
                            raise PermissionError('unsupported_format')
                        raw = path.read_bytes()
                        digest = hashlib.sha256(raw).hexdigest()
                        sidecar_hashes = self._sidecar(path, root)[1] if suffix in IMAGE_FORMATS else []
                        fingerprint = hashlib.sha256(json.dumps([digest, sidecar_hashes, self.settings['source_mappings'], self.settings['products']], sort_keys=True).encode()).hexdigest()
                        old = before.get(file_key)
                        if old and old['fingerprint'] == fingerprint and old['status'] == 'parsed':
                            report['unchanged'] += 1
                        else:
                            parsed, encoding = self._parse(path, root, raw)
                            if not parsed:
                                raise ValueError('no_extractable_content')
                            db.execute('DELETE FROM sources WHERE file_key=?', (file_key,))
                            for position, (text, location, meta) in enumerate(parsed):
                                sid = 'src_' + hashlib.sha256(f'{file_key}\0{digest}\0{position}'.encode()).hexdigest()[:24]
                                product_id, version = self._product(file_key, sid)
                                location.update(relative_path=relative, absolute_path=str(actual))
                                meta.update(encoding=encoding, encoding_fallback=encoding not in ('utf-8', 'utf-8-bom', 'binary:image'),
                                            trust='untrusted', version=version, scope='general' if product_id == 'general' else 'product_specific')
                                meta.setdefault('role', 'unknown')
                                meta.setdefault('date', None)
                                meta.setdefault('conversation_id', None)
                                snippet = redact(text)
                                # Bound any individual model-facing record, preserving coverage explicitly.
                                meta['truncated'] = len(snippet) > 4000
                                db.execute('INSERT INTO sources VALUES(?,?,?,?,?,?,?,?)', (sid, file_key, kind, product_id, digest,
                                           json.dumps(location, ensure_ascii=False), snippet[:4000], json.dumps(_safe_metadata(meta), ensure_ascii=False)))
                            report['modified' if old else 'added'] += 1
                    except PermissionError as exc:
                        status, reason = 'excluded', str(exc)
                    except ImportError as exc:
                        status, reason = 'failed', str(exc)
                    except (OSError, ValueError, TypeError, KeyError, csv.Error) as exc:
                        status, reason = 'failed', ('parse_error:' + type(exc).__name__)
                        if isinstance(exc, ValueError) and str(exc).split(':')[0] in ('decode_failed', 'pdf_no_extractable_text', 'no_extractable_content', 'image_extension_mismatch', 'product_mapping_conflict'):
                            reason = str(exc).split(':')[0]
                    if status != 'parsed':
                        db.execute('DELETE FROM sources WHERE file_key=?', (file_key,))
                        report['issues'].append({'library_type': kind, 'relative_path': relative, 'reason': reason, 'status': status})
                    report[status] += 1
                    counts[status] += 1
                    db.execute('INSERT OR REPLACE INTO files VALUES(?,?,?,?,?,?,?)', (file_key, kind, relative, digest, fingerprint, status, reason))
                for deleted in set(before) - seen:
                    db.execute('DELETE FROM sources WHERE file_key=?', (deleted,))
                    db.execute('DELETE FROM files WHERE file_key=?', (deleted,))
                    report['deleted'] += 1
                report['by_library'][kind] = counts
        report['coverage'] = report['parsed'] / report['files_seen'] if report['files_seen'] else None
        report['updated_at'] = datetime.now(timezone.utc).isoformat()
        return report

    @staticmethod
    def _row(row):
        value = dict(row)
        value.pop('file_key', None)
        value['location'] = json.loads(value['location'])
        value['metadata'] = json.loads(value['metadata'])
        return value

    def get(self, source_id):
        with self._connect() as db:
            row = db.execute('SELECT * FROM sources WHERE source_id=?', (source_id,)).fetchone()
        if row is None:
            raise ValueError('来源不存在或已失效：' + str(source_id))
        return self._row(row)

    def search(self, query, library_type=None, product_id=None, limit=10, **filters):
        if type(limit) is not int or limit < 1 or limit > 1000:
            raise ValueError('limit 必须为 1 到 1000 的整数')
        if library_type is not None and library_type not in LIBRARY_TYPES:
            raise ValueError('未知资料库类型')
        supported = {'role', 'date_from', 'date_to', 'conversation_id', 'subdirectory', 'version', 'is_image'}
        if set(filters) - supported:
            raise ValueError('未知检索筛选条件：' + ','.join(set(filters) - supported))
        if filters.get('subdirectory') and any(p == '..' for p in re.split(r'[/\\]', filters['subdirectory'])):
            raise ValueError('检索子目录不得越界')
        clauses, params = [], []
        for column, value in (('library_type', library_type), ('product_id', product_id)):
            if value is not None:
                clauses.append(column + '=?')
                params.append(value)
        with self._connect() as db:
            rows = db.execute('SELECT * FROM sources' + (' WHERE ' + ' AND '.join(clauses) if clauses else ''), params).fetchall()
        query_grams = _ngrams(str(query))
        found = []
        for raw in rows:
            row = self._row(raw)
            meta = row['metadata']
            if any(meta.get(k) != filters[k] for k in ('role', 'conversation_id', 'version', 'is_image') if k in filters):
                continue
            date = meta.get('date')
            if filters.get('date_from') and (not isinstance(date, str) or date < filters['date_from']):
                continue
            if filters.get('date_to') and (not isinstance(date, str) or date[:10] > filters['date_to'][:10]):
                continue
            directory = filters.get('subdirectory')
            if directory and not row['location']['relative_path'].startswith(directory.replace('\\', '/').rstrip('/') + '/'):
                continue
            grams = _ngrams(row['snippet'])
            intersection = query_grams & grams
            if query_grams and not intersection:
                continue
            row['score'] = sum(2 if len(g) > 1 else 0.25 for g in intersection) / max(1, len(query_grams))
            if query and str(query).lower() in row['snippet'].lower():
                row['score'] += 2
            found.append(row)
        return sorted(found, key=lambda r: (-r['score'], r['source_id']))[:limit]

    def chat_statistics(self, query='', product_id=None, **filters):
        """Count only explicitly customer-authored matched fragments; no customer claims.

        Counts apply to the bounded returned search scope, not an inferred population.
        An explicit conversation ID is namespaced by source file to prevent accidental
        joins between independent exports using the same local session numbering.
        """
        filters.pop('role', None)
        rows = self.search(query, library_type='chat', product_id=product_id, limit=1000,
                           role='customer', **filters)
        conversations = {(r['location']['relative_path'], str(r['metadata']['conversation_id'])) for r in rows
                         if r['metadata']['conversation_id'] is not None}
        return {'matched_fragments': len(rows), 'explicit_conversations': len(conversations),
                'unknown_conversation_fragments': sum(r['metadata']['conversation_id'] is None for r in rows),
                'distinct_snippets': len({r['snippet'] for r in rows}), 'customer_count': None,
                'percentage': None, 'scope': 'matched_customer_fragments',
                'limit_reached': len(rows) == 1000,
                'deduplication': '原始命中按来源位置；重复片段按脱敏全文；会话按文件及原文明示会话ID；无可靠身份不统计客户数。'}
